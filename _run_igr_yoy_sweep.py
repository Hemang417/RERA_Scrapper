import sys
import time
import igr_maharashtra_search as igr
from openpyxl import Workbook
from openpyxl.styles import Font

YEARS = list(range(2020, 2027))
REGION = "mumbai"
DISTRICT = "mumbai"
VILLAGE = "Parel"
PROPERTY_NUMBER = 10

all_rows = []
summary = []
sweep_started_at = time.time()

for year in YEARS:
    year_started_at = time.time()
    print(f"\n===== YEAR {year} =====", flush=True)
    result = igr.search_by_property(REGION, DISTRICT, VILLAGE, PROPERTY_NUMBER, year=year, fetch_index_ii=True)
    found = result.get("found")
    rows = result.get("rows", [])
    note = result.get("note", "")

    # CONFIRMED LIVE 2026-10-15: even with the field-reset bug fixed (every
    # field verified correct before AND after submission), the identical
    # search has still returned a genuine stable "no result" on one run and
    # a real match on another -- residual portal-side backend flakiness,
    # not a code bug. Up to 2 extra retries (3 attempts total) on an empty
    # reading before trusting it as a true negative.
    attempts_used = 1
    while not rows and attempts_used < 3:
        attempts_used += 1
        print(f"[INFO] Year {year}: no rows on attempt {attempts_used - 1} -- retrying "
              f"(attempt {attempts_used}/3) to rule out portal flakiness.")
        result_n = igr.search_by_property(REGION, DISTRICT, VILLAGE, PROPERTY_NUMBER, year=year, fetch_index_ii=True)
        if result_n.get("rows"):
            found = result_n.get("found")
            rows = result_n.get("rows", [])
            note = result_n.get("note", "")

    year_elapsed = time.time() - year_started_at
    print(f"[INFO] Year {year}: found={found}, rows={len(rows)} (after {attempts_used} attempt(s), "
          f"{year_elapsed:.1f}s)")
    summary.append({
        "year": year, "found": bool(rows), "row_count": len(rows),
        "attempts": attempts_used, "note": note, "elapsed_seconds": round(year_elapsed, 1),
    })
    for row in rows:
        row_copy = dict(row)
        row_copy["search_year"] = year
        all_rows.append(row_copy)

sweep_elapsed = time.time() - sweep_started_at
print(f"\n===== SWEEP COMPLETE in {sweep_elapsed / 60:.1f} min ({sweep_elapsed:.0f}s) =====")
for s in summary:
    print(s)

# -- Build the Excel workbook --------------------------------------------
wb = Workbook()

ws_summary = wb.active
ws_summary.title = "YoY Summary"
ws_summary.append(["Year", "Documents Found", "Attempts Needed", "Time Taken (s)", "Note"])
for cell in ws_summary[1]:
    cell.font = Font(bold=True)
for s in summary:
    ws_summary.append([s["year"], s["row_count"], s["attempts"], s["elapsed_seconds"], s["note"]])
ws_summary.append(["TOTAL", sum(s["row_count"] for s in summary), "", round(sweep_elapsed, 1), ""])
for cell in ws_summary[ws_summary.max_row]:
    cell.font = Font(bold=True)
ws_summary.column_dimensions["A"].width = 8
ws_summary.column_dimensions["B"].width = 16
ws_summary.column_dimensions["C"].width = 16
ws_summary.column_dimensions["D"].width = 14
ws_summary.column_dimensions["E"].width = 80

ws_docs = wb.create_sheet("Documents")
headers = [
    "Search Year", "DocNo", "DName", "RDate", "SROName",
    "Seller Name", "Purchaser Name", "Property Description", "SROCode",
    "Status", "IndexII", "index_ii_text",
]
ws_docs.append(headers)
for cell in ws_docs[1]:
    cell.font = Font(bold=True)
for row in all_rows:
    ws_docs.append([
        row.get("search_year"),
        row.get("docno"),
        row.get("dname"),
        row.get("rdate"),
        row.get("sroname"),
        row.get("seller name"),
        row.get("purchaser name"),
        row.get("property description"),
        row.get("srocode"),
        row.get("status"),
        row.get("indexii"),
        row.get("index_ii_text"),
    ])
widths = [10, 12, 28, 12, 16, 40, 40, 50, 10, 10, 10, 80]
for i, w in enumerate(widths, start=1):
    ws_docs.column_dimensions[ws_docs.cell(row=1, column=i).column_letter].width = w

out_path = "IGR_Property_YoY_2020-2026.xlsx"
wb.save(out_path)
print(f"\n[INFO] Wrote {out_path} ({len(all_rows)} total document rows across {len(YEARS)} years)")
