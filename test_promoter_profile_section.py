"""
Tests for _promoter_trust_signals and _append_promoter_profile_section --
the consolidated Promoter Profile section (see CLAUDE.md's "2b. Identity
and group passes" and company_charter.py's own docstring on why this does
NOT re-render Group Companies' table/diagram, and why it states rather than
estimates personal-wealth data that isn't publicly available).

Run directly: python test_promoter_profile_section.py
"""

import json
import os
import shutil
from unittest import mock

import docx

import company_charter as cc
import states

_SCRATCH = os.path.join("output", "company_charters", "_test_scratch_promoter_profile")
_PRANAMI_FACTS = os.path.join("output", "company_charters", "Company_Charter_Pranami_Bliss_P51800077150.facts.json")


# --- _promoter_trust_signals: pure-function checks --------------------------

def test_no_cin_means_not_checked_not_clean():
    """Absence of a checkable identifier must never read as a clean result
    -- see guardrails.md's "no finding never means no check" discipline."""
    signals = cc._promoter_trust_signals({})
    by_name = {s["signal"]: s for s in signals}
    assert by_name["MCA Charge Transparency"]["status"] == "Not checked"
    assert by_name["Credit Rating Coverage"]["status"] == "Not checked"
    assert by_name["Insolvency (IBBI)"]["status"] == "Not checked"
    assert by_name["Group Entity Transparency"]["status"] == "Not checked"
    print("test_no_cin_means_not_checked_not_clean: PASS")


def test_open_charge_is_flagged_with_lender_and_amount():
    # "amount" is a raw scraped STRING (sometimes with commas, sometimes
    # unparseable) -- never an int/float -- see summarise_charges's own
    # docstring on why this matters.
    facts = {"company_profile_check": {"charges": [
        {"is_open": True, "charge_holder": "HDFC Bank", "amount": "90,300,000"},
        {"is_open": False, "charge_holder": "Axis Bank", "amount": "1,000,000"},
    ]}}
    signals = cc._promoter_trust_signals(facts)
    row = next(s for s in signals if s["signal"] == "MCA Charge Transparency")
    assert row["status"] == "Flagged"
    assert "HDFC Bank" in row["detail"]
    assert "Axis Bank" not in row["detail"]  # closed charge must not count as open
    print("test_open_charge_is_flagged_with_lender_and_amount: PASS")


def test_unparseable_amount_says_so_rather_than_a_misleading_zero():
    facts = {"company_profile_check": {"charges": [
        {"is_open": True, "charge_holder": "HDFC Bank", "amount": "not disclosed"},
    ]}}
    row = next(s for s in cc._promoter_trust_signals(facts) if s["signal"] == "MCA Charge Transparency")
    assert row["status"] == "Flagged"
    assert "unreadable" in row["detail"]
    print("test_unparseable_amount_says_so_rather_than_a_misleading_zero: PASS")


def test_no_open_charges_is_clean_not_silent():
    facts = {"company_profile_check": {"charges": [{"is_open": False, "charge_holder": "Axis Bank", "amount": "1"}]}}
    row = next(s for s in cc._promoter_trust_signals(facts) if s["signal"] == "MCA Charge Transparency")
    assert row["status"] == "Clean"
    print("test_no_open_charges_is_clean_not_silent: PASS")


def test_unrated_promoter_is_not_applicable_not_flagged():
    """No public rating is ordinary for a private, unlisted promoter -- must
    not be presented as a red flag."""
    facts = {"credit_rating_check": {"promoter": {"ratings": [], "not_found_agencies": ["CRISIL", "ICRA"]}}}
    row = next(s for s in cc._promoter_trust_signals(facts) if s["signal"] == "Credit Rating Coverage")
    assert row["status"] == "Not applicable"
    print("test_unrated_promoter_is_not_applicable_not_flagged: PASS")


def test_ibbi_found_process_is_flagged():
    facts = {"ibbi_insolvency_check": {"found_process": True}}
    row = next(s for s in cc._promoter_trust_signals(facts) if s["signal"] == "Insolvency (IBBI)")
    assert row["status"] == "Flagged"
    print("test_ibbi_found_process_is_flagged: PASS")


def test_bombay_hc_declined_consent_is_not_checked():
    facts = {"bombay_hc_check": {"attempted": False, "note": "Declined this pass."}}
    row = next(s for s in cc._promoter_trust_signals(facts) if s["signal"] == "Bombay High Court Case Status")
    assert row["status"] == "Not checked"
    print("test_bombay_hc_declined_consent_is_not_checked: PASS")


def test_igr_row_absent_for_a_non_maharashtra_promoter():
    """_promoter_trust_signals reads the ACTIVE render profile
    (_state_profile()), not facts["state"] directly -- a rendering
    decision, so a caller re-rendering the same saved facts under an
    explicit different profile (see test_state_labels.py) must get that
    profile's answer, not the original run's. Set here the same way
    _fill_template_inner sets it around its own render."""
    gj_profile = states.get_profile("GJ")
    with mock.patch.object(cc, "_ACTIVE_STATE_PROFILE", gj_profile):
        signals = cc._promoter_trust_signals({})
    assert not any("IGR Maharashtra" in s["signal"] for s in signals)
    print("test_igr_row_absent_for_a_non_maharashtra_promoter: PASS")


def test_igr_row_present_when_no_profile_is_active_defaults_to_mh():
    """No _ACTIVE_STATE_PROFILE set (e.g. _promoter_trust_signals called
    outside a _fill_template render) -- _state_profile() falls back to
    Maharashtra, same convention as every other state-aware render here."""
    with mock.patch.object(cc, "_ACTIVE_STATE_PROFILE", None):
        signals = cc._promoter_trust_signals({})
    assert any("IGR Maharashtra" in s["signal"] for s in signals)
    print("test_igr_row_present_when_no_profile_is_active_defaults_to_mh: PASS")


def test_group_entity_transparency_reports_undisclosed_relationships():
    facts = {"group_companies_check": {"found": True, "companies": [{"name": "X", "cin": "Y", "basis": []}],
                                        "undisclosed_relationship_counts": {"related": 3}}}
    row = next(s for s in cc._promoter_trust_signals(facts) if s["signal"] == "Group Entity Transparency")
    assert "3 further relationship" in row["detail"]
    print("test_group_entity_transparency_reports_undisclosed_relationships: PASS")


# --- _append_promoter_profile_section: rendering ----------------------------

def _facts():
    with open(_PRANAMI_FACTS, encoding="utf-8") as f:
        return json.load(f)


def _build(facts: dict, variant: str) -> docx.Document:
    os.makedirs(_SCRATCH, exist_ok=True)
    out = os.path.join(_SCRATCH, f"{variant}.docx")
    cc._fill_template("P51800077150", facts, out, doc_variant=variant)
    return docx.Document(out)


def test_section_renders_once_and_does_not_duplicate_group_companies_table():
    facts = _facts()
    facts["promoter_external_research"] = {"summary": "A test brief profile sentence."}
    d = _build(facts, "internal")
    headings = [p.text for p in d.paragraphs if p.text.strip() == "Promoter Profile" or "Promoter Profile" in p.text]
    assert sum(1 for p in d.paragraphs if p.text.strip().endswith("Promoter Profile")) == 1, headings

    # The company-name/CIN/basis table lives ONLY in Group / Affiliated
    # Companies -- Promoter Profile must reference it, never reproduce it.
    company_col_tables = [
        t for t in d.tables
        if t.rows and [c.text for c in t.rows[0].cells][:3] == ["Company Name", "CIN / Identifier", "Basis for Link"]
    ]
    assert len(company_col_tables) == 1, "the company/CIN/basis table must render exactly once, in Group Companies"
    shutil.rmtree(_SCRATCH, ignore_errors=True)
    print("test_section_renders_once_and_does_not_duplicate_group_companies_table: PASS")


def test_not_publicly_available_block_is_present_and_honest():
    facts = _facts()
    d = _build(facts, "internal")
    full_text = "\n".join(p.text for p in d.paragraphs)
    assert "net worth" in full_text.lower()
    assert "not available from any public source" in full_text.lower() or "not estimated here" in full_text.lower()
    shutil.rmtree(_SCRATCH, ignore_errors=True)
    print("test_not_publicly_available_block_is_present_and_honest: PASS")


def test_section_is_silent_when_no_identity_data_at_all():
    """_fill_template needs a schema-complete facts dict (see
    test_facts_normalization.py's own _facts()-based fixtures) -- a
    minimal ad-hoc dict isn't a valid input on its own, so this strips
    just the promoter-identity keys from a real, complete fixture instead
    of building a fake one from scratch."""
    facts = _facts()
    facts["corporate_identity"] = {}
    facts["company_profile_check"] = {"found": False}
    d = _build(facts, "internal")
    assert not any(p.text.strip().endswith("Promoter Profile") for p in d.paragraphs)
    shutil.rmtree(_SCRATCH, ignore_errors=True)
    print("test_section_is_silent_when_no_identity_data_at_all: PASS")


if __name__ == "__main__":
    test_no_cin_means_not_checked_not_clean()
    test_open_charge_is_flagged_with_lender_and_amount()
    test_unparseable_amount_says_so_rather_than_a_misleading_zero()
    test_no_open_charges_is_clean_not_silent()
    test_unrated_promoter_is_not_applicable_not_flagged()
    test_ibbi_found_process_is_flagged()
    test_igr_row_absent_for_a_non_maharashtra_promoter()
    test_igr_row_present_when_no_profile_is_active_defaults_to_mh()
    test_bombay_hc_declined_consent_is_not_checked()
    test_group_entity_transparency_reports_undisclosed_relationships()
    test_section_renders_once_and_does_not_duplicate_group_companies_table()
    test_not_publicly_available_block_is_present_and_honest()
    test_section_is_silent_when_no_identity_data_at_all()
    print("\nAll tests passed.")
